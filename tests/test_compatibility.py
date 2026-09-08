import contextlib
import hashlib
import importlib
import io
import json
from pathlib import Path
import sys
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import MagicMock, patch

import openpyxl

from src import excel_access, platform_utils
from src.gui.SpExGui import SpExGui
from src.gui.ToggleGui import ToggleGui
from src.write.ecount import EcountWriter as ecount_writer_module


ROOT = Path(__file__).resolve().parents[1]


class FileInputTests(unittest.TestCase):
    def test_both_file_choosers_accept_native_unicode_paths(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / '출고 파일.xlsx'
            workbook = openpyxl.Workbook()
            workbook.active.title = '주문'
            workbook.save(path)
            workbook.close()
            for gui_class in (SpExGui, ToggleGui):
                with self.subTest(gui=gui_class.__name__):
                    gui = SimpleNamespace(fileNameLabel=MagicMock(), sheetCombobox=MagicMock())
                    gui_class.setSpExFileName(gui, str(path))
                    gui.sheetCombobox.configure.assert_called_once_with(values=['주문'])
            # On Windows this also detects unclosed input handles.
            path.unlink()

    def test_cancel_preserves_current_file(self):
        for gui_class in (SpExGui, ToggleGui):
            with self.subTest(gui=gui_class.__name__), patch('tkinter.filedialog.askopenfilename', return_value=''):
                gui = SimpleNamespace(setSpExFileName=MagicMock(), fileNameLabel=MagicMock(), spExFile='기존 파일')
                gui_class.findFile(gui)
                gui.setSpExFileName.assert_not_called()

    def test_file_mode_does_not_import_excel_automation(self):
        self.assertNotIn('xlwings', sys.modules)
        self.assertTrue(excel_access.sheet_names(ROOT / 'example.xlsx'))
        self.assertNotIn('xlwings', sys.modules)


class DesktopOperationTests(unittest.TestCase):
    def test_document_opener_on_both_platforms(self):
        path = str((ROOT / '한글 파일 & $(example).xlsx').resolve())
        with patch.object(platform_utils.sys, 'platform', 'darwin'), patch.object(platform_utils.subprocess, 'run') as run:
            platform_utils.open_file(path)
            run.assert_called_once_with(['/usr/bin/open', path], check=True)
        with patch.object(platform_utils.sys, 'platform', 'win32'), patch.object(platform_utils.os, 'startfile', create=True) as startfile:
            platform_utils.open_file(path)
            startfile.assert_called_once_with(path)

    def test_missing_excel_has_an_actionable_error(self):
        fake_xlwings = SimpleNamespace(apps=SimpleNamespace(active=None))
        with patch.dict(sys.modules, {'xlwings': fake_xlwings}):
            with self.assertRaisesRegex(ValueError, 'Excel을 실행'):
                excel_access.active_book()

    def test_snapshot_uses_source_instance_and_releases_temporary_files(self):
        book = MagicMock()
        snapshot = book.app.books.add.return_value

        def save(path):
            workbook = openpyxl.Workbook()
            workbook.active.title = excel_access.SNAPSHOT_SHEET
            workbook.active.append(['한글', 42])
            workbook.save(path)
            workbook.close()

        snapshot.save.side_effect = save
        with patch.object(excel_access, 'active_book', return_value=book):
            with excel_access.active_sheet() as sheet:
                self.assertEqual(next(sheet.values), ('한글', 42))
                path = Path(snapshot.save.call_args.args[0])
                self.assertEqual(path.suffix, '.xlsx')
                snapshot.close.assert_called_once()
                self.assertTrue(path.is_file())
            self.assertFalse(path.parent.exists())
        book.sheets.active.copy.assert_called_once_with(
            before=snapshot.sheets[0], name=excel_access.SNAPSHOT_SHEET)
        book.activate.assert_called_once()
        book.close.assert_not_called()
        book.save.assert_not_called()

    def test_failed_snapshot_closes_only_the_copy(self):
        book = MagicMock()
        snapshot = book.app.books.add.return_value
        snapshot.save.side_effect = OSError('cannot save')
        with patch.object(excel_access, 'active_book', return_value=book):
            with self.assertRaisesRegex(OSError, 'cannot save'):
                with excel_access.active_sheet():
                    self.fail('A failed snapshot must not be exposed')
        snapshot.close.assert_called_once()
        book.activate.assert_called_once()
        book.close.assert_not_called()
        self.assertFalse(Path(snapshot.save.call_args.args[0]).parent.exists())

    def test_output_open_error_survives_cleanup_with_partial_snapshot_reader(self):
        book = MagicMock()
        snapshot = book.app.books.add.return_value
        snapshot_path = None

        def save(path):
            nonlocal snapshot_path
            snapshot_path = Path(path)
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = excel_access.SNAPSHOT_SHEET
            sheet.append(['시간', '챙길것', '품목', '수량', '보험사', '지점',
                          '주문자이름', '주소', '전화번호', '단가', '금액', '수금'])
            sheet.append(['#끝'])
            for row in range(1000):
                sheet.append([f'unread trailing row {row} ' + ('x' * 200)])
            workbook.save(path)
            workbook.close()

        snapshot.save.side_effect = save
        opener_error = RuntimeError('output opener failed')
        caught = None
        writer = None
        with TemporaryDirectory() as output_directory, \
                patch.object(excel_access, 'active_book', return_value=book), \
                patch.object(ecount_writer_module, 'getTempDir', return_value=output_directory), \
                patch.object(ecount_writer_module, 'open_file', side_effect=opener_error):
            try:
                with excel_access.active_sheet() as sheet:
                    writer = ecount_writer_module.EcountWriter.fromSheet(sheet)
                    writer.getDocsFromSpEx()
            except Exception as error:
                caught = error

        self.assertIs(caught, opener_error)
        self.assertIsNotNone(caught.__traceback__)
        self.assertIsNotNone(writer.spExReader.it.gi_frame)
        self.assertIsNotNone(snapshot_path)
        self.assertFalse(snapshot_path.parent.exists())

    def test_selected_range_preserves_headers_and_rows(self):
        book = MagicMock()
        book.sheets.active.range.return_value.value = ['품목', '수량']
        book.app.selection.value = [['상품1', 2], ['상품2', 3]]
        with patch.object(excel_access, 'active_book', return_value=book):
            self.assertEqual(excel_access.selected_range(), [['품목', '수량'], ['상품1', 2], ['상품2', 3]])


class ConversionRegressionTests(unittest.TestCase):
    def test_sample_outputs_match_before_the_port(self):
        # Digests of every saved cell, captured from master (44c85da).
        # Ignore ZIP timestamps and XML metadata; preserve actual document content.
        cases = [
            ('src.write.ecount.EcountWriter', 'EcountWriter',
             'aef7bd4e5a4f65fd0a37496bfae8bdb6503e279c3f09a924a08c6c51dbcedafc'),
            ('src.write.wehago.WehagoWriter', 'WehagoWriter',
             '1a333566a5ba16384392807e747adf87c8296993500446f3df3ce5d7fd87ab37'),
        ]
        for module_name, class_name, expected in cases:
            with self.subTest(writer=class_name), TemporaryDirectory() as directory:
                module = importlib.import_module(module_name)
                workbook = openpyxl.load_workbook(ROOT / 'example.xlsx', read_only=True, data_only=True)
                try:
                    with patch.object(module, 'getTempDir', return_value=directory), \
                            patch.object(module, 'open_file') as opened, \
                            contextlib.redirect_stdout(io.StringIO()):
                        getattr(module, class_name).fromSheet(workbook.active).getDocsFromSpEx()
                finally:
                    workbook.close()
                paths = list(Path(directory).glob('*.xlsx'))
                self.assertEqual(len(paths), 1)
                opened.assert_called_once()
                self.assertEqual(Path(opened.call_args.args[0]), paths[0])
                output = openpyxl.load_workbook(paths[0], read_only=True, data_only=True)
                try:
                    values = {sheet.title: list(sheet.values) for sheet in output}
                finally:
                    output.close()
                actual = hashlib.sha256(json.dumps(values, ensure_ascii=False, default=str).encode()).hexdigest()
                self.assertEqual(actual, expected)


if __name__ == '__main__':
    unittest.main()
