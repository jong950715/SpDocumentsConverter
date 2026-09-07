"""Read saved workbooks or take a temporary snapshot of desktop Excel."""

from contextlib import contextmanager
from pathlib import Path
from tempfile import TemporaryDirectory

import openpyxl


SNAPSHOT_SHEET = 'ConverterInput'


def sheet_names(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return workbook.sheetnames
    finally:
        workbook.close()


@contextmanager
def open_sheet(path, name):
    if not path or not Path(path).is_file():
        raise ValueError('변환할 Excel 파일을 먼저 선택해주세요.')
    if not name:
        raise ValueError('변환할 시트를 선택해주세요.')
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        yield workbook[name]
    finally:
        workbook.close()


def active_book():
    # File-based conversion must not initialize Excel automation (or its exit hooks).
    import xlwings

    app = xlwings.apps.active
    book = app.books.active if app is not None else None
    if book is None:
        raise ValueError('Excel을 실행하고 변환할 문서를 열어주세요.')
    return book


@contextmanager
def active_sheet():
    book = active_book()
    sheet = book.sheets.active
    with TemporaryDirectory(prefix='spdocuments-') as directory:
        path = Path(directory) / 'input.xlsx'
        snapshot = None
        try:
            # Use the source book's Excel instance, including on Windows with
            # multiple instances. Always save as xlsx, even for unsaved/xls books.
            snapshot = book.app.books.add()
            sheet.copy(before=snapshot.sheets[0], name=SNAPSHOT_SHEET)
            # Only the disposable copy loses macros when saved as xlsx.
            with book.app.properties(display_alerts=False):
                snapshot.save(str(path))
        finally:
            try:
                if snapshot is not None:
                    snapshot.close()
            finally:
                book.activate()
        with open_sheet(path, SNAPSHOT_SHEET) as saved_sheet:
            yield saved_sheet


def selected_range():
    book = active_book()
    selection = book.app.selection
    if selection is None:
        raise ValueError('Excel에서 변환할 셀 영역을 선택해주세요.')
    data = [book.sheets.active.range('TITLES').value]
    data.extend(selection.value)
    return data
